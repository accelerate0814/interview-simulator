import io
import os
import json
import uuid
import random
import asyncio
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Tuple
from contextlib import asynccontextmanager

import numpy as np
import aiosqlite
from openai import AsyncOpenAI
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, Header, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

load_dotenv()

SKILL_PATH = Path(__file__).parent / "skill.md"
SYSTEM_PROMPT = SKILL_PATH.read_text(encoding="utf-8")
DB_PATH = Path(__file__).parent / "interviews.db"

# ── Resume + JD (定向出题) ───────────────────────
# Text pulled from an uploaded resume / a pasted JD is truncated before it goes
# into the system prompt so one huge file can't blow the context window.
RESUME_MAX_CHARS = 12000
JD_MAX_CHARS = 8000

REPORT_PROMPT = """\
你是专业的面试报告生成助手。根据面试对话，生成结构化报告。

只返回如下JSON，不要有任何额外文字：
{
  "role": "面试岗位（从对话提取）",
  "level": "经验层级",
  "overall_score": 7.5,
  "verdict": "Hire",
  "questions": [
    {
      "question": "面试题目完整表述",
      "user_answer": "候选人回答摘要（2-3句）",
      "standard_answer": "完整标准答案，覆盖所有关键知识点，不少于200字",
      "score": 8,
      "feedback": "对本次回答的具体点评（1-2句）"
    }
  ],
  "strengths": ["优势1", "优势2", "优势3"],
  "improvements": ["改进点1", "改进点2", "改进点3"],
  "recommended_topics": ["建议深入学习的主题1", "主题2", "主题3"]
}

注意：
- questions 只含真正的面试题，不含开场白和结束语
- verdict 只能是：Strong Hire / Hire / Lean Hire / Lean No Hire / No Hire
- overall_score 为 0-10 小数"""

# ── Cross-session question de-dup ────────────────
# History small (<= DEDUP_FULL_LIMIT): inline every past question.
# History large: pick the questions most *semantically related* to the current
# interview via embeddings + in-memory cosine (RETRIEVAL_TOP_K of them), plus a
# per-category tally. If embeddings aren't available, degrade to "most recent
# DEDUP_RECENT_N" — the old behaviour.
DEDUP_FULL_LIMIT = 50
DEDUP_RECENT_N = 30
RETRIEVAL_TOP_K = 20
RETRIEVAL_MIN_EMBEDDED = 5  # need at least this many embedded rows to bother searching

# Commands (see skill.md "Special Commands") produce hints / explanations /
# scorecards, not new interview questions — skip question recording for those.
INTERVIEW_COMMANDS = {
    "hint", "skip", "explain", "score", "harder", "easier",
    "end", "restart", "switch",
}

EXTRACT_PROMPT = """你是面试内容分析器。从下面这段“面试官”的发言里，提取其中**正式提出的面试题目**。

排除以下内容，不要当作题目：
- 寒暄、开场白、询问岗位/级别/时长等信息收集
- 对候选人上一题回答的点评、打分、总结
- “还有要补充的吗”这类追问
- 命令回复（提示 / 解析 / 当前得分等）

只返回 JSON，不要任何多余文字：
{"questions":[{"text":"题目的完整、可独立理解的表述（去掉“第1题”之类的编号前缀）","category":"类别"}]}

category 从下面选一个最接近的：算法、系统设计、编程语言/框架、数据库、网络、操作系统、项目经验、行为面试、案例分析、专业知识、其他。

如果这段发言没有提出任何新的正式面试题，返回：{"questions":[]}"""

# ── 面试知识库（用户导入的真实面试题）──────────────
KB_MAX_IMPORT_CHARS = 20000          # cap the text handed to the extractor
KB_STYLE_EXAMPLE_N = 5              # few-shot examples pulled for "style" turns
# Per-turn branch weights for how the next question is produced. Tunable.
KB_REPLAY_PROB = 0.20              # ask a real question from the bank verbatim
KB_STYLE_PROB = 0.50              # generate a new question in the bank's style
# remaining 0.30 -> free generation (existing behaviour, no block added)

KB_EXTRACT_PROMPT = """你是面试题目结构化助手。用户给你一整篇他自己整理的面试笔记（markdown 或 word 转出的纯文本），里面记录了他真实面试遇到过的题目。请把其中的**面试题目**逐条拆出来。

要求：
- 一条题目一个对象，题目文本尽量保留原文表述，去掉“第1题”“Q3”之类的编号前缀
- category：从下面选一个最接近的，识别不出就用 null —— 算法、系统设计、编程语言/框架、数据库、网络、操作系统、项目经验、行为面试、案例分析、专业知识、其他
- company_or_role：如果笔记里能看出这条题目对应的公司或岗位（如“字节-后端”“腾讯 二面”），填进去；看不出就 null
- 只提取真正的面试问题，跳过纯笔记、答案、心得、时间线

只返回 JSON，不要任何多余文字：
{"questions":[{"question":"题目完整文本","category":"类别或null","company_or_role":"公司/岗位或null"}]}

如果整篇没有任何可识别的面试题，返回：{"questions":[]}"""

# Persistent, cross-session memory of *who the candidate is* (name, target role,
# level, years, stack, company, stated preferences). Re-derived after each turn
# from the running profile + the newest conversation, then upserted.
PROFILE_MAX_CHARS = 1500
PROFILE_PROMPT = f"""你在维护一份"面试候选人档案"，用于**跨场次记忆**同一个人。
给你【当前档案】和【最新对话片段】，输出合并更新后的完整档案。

只记录**稳定、可跨场次复用**的事实，例如：
- 姓名 / 希望被怎么称呼
- 目标岗位、级别（intern/junior/mid/senior/staff…）
- 工作年限、技术栈 / 专业方向、当前或过往公司、学历
- 候选人明确表达的偏好或约束（如"用中文面试""重点练系统设计""明天有面试"）

不要记录：某道题的作答、面试官的提问、分数、临时的一句话。

规则：
- 合并新事实、保留旧事实；只有被明确推翻时才修改或删除。
- 每条一行，格式 `- 字段: 值`；总长度不超过 {PROFILE_MAX_CHARS} 字。
- 没有任何可记录的稳定事实时，返回空字符串。

只返回 JSON：{{"profile": "更新后的完整档案文本"}}"""

# ── Default AI client (from .env) ───────────────
_default_ai = AsyncOpenAI(
    api_key=os.getenv("DEEPSEEK_API_KEY", ""),
    base_url="https://api.deepseek.com",
)
_default_model = "deepseek-chat"

# ── Embedding client (separate from the chat provider) ───────────────
# The chat providers people plug in here (DeepSeek, Moonshot, OpenRouter, ...)
# mostly have NO embeddings endpoint, so embeddings use their own config.
# Set EMBEDDING_API_KEY in backend/.env to turn semantic retrieval on; without
# it the de-dup logic just falls back to the recency-based behaviour.
_embed_api_key = os.getenv("EMBEDDING_API_KEY", "")
_embed_base_url = os.getenv("EMBEDDING_BASE_URL", "https://api.openai.com/v1")
_embed_model = os.getenv("EMBEDDING_MODEL", "text-embedding-3-small")
_embed_client = (
    AsyncOpenAI(api_key=_embed_api_key, base_url=_embed_base_url)
    if _embed_api_key
    else None
)
EMBEDDINGS_ENABLED = _embed_client is not None


# ── DB ───────────────────────────────────────────
async def init_db():
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute("""
            CREATE TABLE IF NOT EXISTS sessions (
                id          TEXT    PRIMARY KEY,
                title       TEXT    NOT NULL DEFAULT '新面试',
                turn_count  INTEGER NOT NULL DEFAULT 0,
                duration_s  INTEGER NOT NULL DEFAULT 0,
                report_json TEXT,
                ended       INTEGER NOT NULL DEFAULT 0,
                created_at  TEXT    NOT NULL DEFAULT (datetime('now','localtime')),
                updated_at  TEXT    NOT NULL DEFAULT (datetime('now','localtime'))
            )
        """)
        await db.execute("""
            CREATE TABLE IF NOT EXISTS messages (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                session_id  TEXT    NOT NULL,
                role        TEXT    NOT NULL,
                content     TEXT    NOT NULL,
                created_at  TEXT    NOT NULL DEFAULT (datetime('now','localtime'))
            )
        """)
        await db.execute(
            "CREATE INDEX IF NOT EXISTS idx_msg_session ON messages(session_id)"
        )
        # Cross-session record of every interview question a given client has
        # already been asked, so future sessions can avoid repeating them.
        await db.execute("""
            CREATE TABLE IF NOT EXISTS asked_questions (
                id            INTEGER PRIMARY KEY AUTOINCREMENT,
                client_id     TEXT    NOT NULL,
                question_text TEXT    NOT NULL,
                category      TEXT,
                session_id    TEXT,
                embedding     TEXT,
                created_at    TEXT    NOT NULL DEFAULT (datetime('now','localtime'))
            )
        """)
        await db.execute(
            "CREATE INDEX IF NOT EXISTS idx_asked_client ON asked_questions(client_id)"
        )
        # One evolving free-text profile per client (name, target role, level,
        # stack, ...) so a brand-new session still "remembers" the candidate.
        await db.execute("""
            CREATE TABLE IF NOT EXISTS client_profile (
                client_id  TEXT PRIMARY KEY,
                profile    TEXT NOT NULL DEFAULT '',
                updated_at TEXT NOT NULL DEFAULT (datetime('now','localtime'))
            )
        """)
        # The candidate's résumé (one per client — only the newest is kept) and
        # any number of pasted job descriptions, used to tailor the questions.
        await db.execute("""
            CREATE TABLE IF NOT EXISTS user_profile (
                client_id       TEXT PRIMARY KEY,
                resume_text     TEXT NOT NULL DEFAULT '',
                resume_filename TEXT,
                updated_at      TEXT NOT NULL DEFAULT (datetime('now','localtime'))
            )
        """)
        await db.execute("""
            CREATE TABLE IF NOT EXISTS job_descriptions (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                client_id   TEXT    NOT NULL,
                title       TEXT    NOT NULL DEFAULT '',
                jd_text     TEXT    NOT NULL,
                created_at  TEXT    NOT NULL DEFAULT (datetime('now','localtime'))
            )
        """)
        await db.execute(
            "CREATE INDEX IF NOT EXISTS idx_jd_client ON job_descriptions(client_id)"
        )
        # 面试知识库：用户导入的真实面试题（全局，不按 client 分）。出题逻辑会
        # 从这里原样重放或抽样做 few-shot 风格参考。
        await db.execute("""
            CREATE TABLE IF NOT EXISTS question_bank (
                id              INTEGER PRIMARY KEY AUTOINCREMENT,
                question_text   TEXT    NOT NULL,
                category        TEXT,
                source_file     TEXT,
                company_or_role TEXT,
                created_at      TEXT    NOT NULL DEFAULT (datetime('now','localtime'))
            )
        """)
        await db.execute(
            "CREATE INDEX IF NOT EXISTS idx_qbank_category ON question_bank(category)"
        )
        # migrations: add columns when upgrading from an older schema. Old rows
        # keep NULL (embedding is backfilled lazily / left empty — that's fine).
        for stmt in (
            "ALTER TABLE sessions ADD COLUMN report_json TEXT",
            "ALTER TABLE sessions ADD COLUMN ended INTEGER NOT NULL DEFAULT 0",
            "ALTER TABLE asked_questions ADD COLUMN embedding TEXT",
            "ALTER TABLE asked_questions ADD COLUMN source TEXT",
        ):
            try:
                await db.execute(stmt)
            except Exception:
                pass
        await db.commit()


@asynccontextmanager
async def lifespan(app: FastAPI):
    await init_db()
    yield


app = FastAPI(title="Interview Simulator", lifespan=lifespan)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


# ── Pydantic models ──────────────────────────────
class LLMConfig(BaseModel):
    api_key: str
    base_url: str
    model: str


class Message(BaseModel):
    role: str
    content: str


class ChatRequest(BaseModel):
    messages: List[Message]
    llm: Optional[LLMConfig] = None
    session_id: Optional[str] = None  # optional: link recorded questions to a session
    use_resume: bool = False          # inject the stored résumé into the system prompt
    jd_id: Optional[int] = None       # inject this saved JD into the system prompt


class JDCreate(BaseModel):
    title: str = ""
    jd_text: str


class TurnSave(BaseModel):
    session_id: str
    user_content: str
    assistant_content: str
    turn_count: int
    duration_s: int
    title: Optional[str] = None


class ReportRequest(BaseModel):
    llm: Optional[LLMConfig] = None


# ── Helpers ──────────────────────────────────────
def make_client(cfg: Optional[LLMConfig]):
    """Return (AsyncOpenAI client, model_name) for the given config or defaults."""
    if cfg and cfg.api_key and cfg.base_url and cfg.model:
        return AsyncOpenAI(api_key=cfg.api_key, base_url=cfg.base_url), cfg.model
    return _default_ai, _default_model


# ── Question de-dup helpers ──────────────────────
async def fetch_asked_questions(client_id: str) -> List[dict]:
    """All questions this client has ever been asked, newest first.

    Note: the (large) embedding column is deliberately not selected here — this
    feeds the prompt text / category tally / the /api/asked-questions endpoint.
    """
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            "SELECT question_text, category, created_at FROM asked_questions "
            "WHERE client_id = ? ORDER BY id DESC",
            (client_id,),
        ) as cur:
            rows = await cur.fetchall()
    return [dict(r) for r in rows]


# ── Embeddings + semantic retrieval ──────────────
async def embed_text(text: str) -> Optional[List[float]]:
    """Embed one string. Returns None if embeddings are disabled or the call fails."""
    if not _embed_client or not text.strip():
        return None
    try:
        resp = await _embed_client.embeddings.create(
            model=_embed_model, input=text[:8000]
        )
        return list(resp.data[0].embedding)
    except Exception:
        return None


async def fetch_question_embeddings(client_id: str) -> List[Tuple[str, List[float]]]:
    """(question_text, vector) for this client's questions that have an embedding."""
    async with aiosqlite.connect(DB_PATH) as db:
        async with db.execute(
            "SELECT question_text, embedding FROM asked_questions "
            "WHERE client_id = ? AND embedding IS NOT NULL",
            (client_id,),
        ) as cur:
            rows = await cur.fetchall()
    out: List[Tuple[str, List[float]]] = []
    for text, emb in rows:
        try:
            vec = json.loads(emb)
        except Exception:
            continue
        if isinstance(vec, list) and vec:
            out.append((text, vec))
    return out


def cosine_top_k(
    query: List[float],
    candidates: List[Tuple[str, List[float]]],
    k: int,
) -> List[Tuple[str, float]]:
    """Brute-force cosine similarity in memory (fine for a few thousand rows).

    candidates: [(text, vector)].  Returns [(text, score)] sorted by score desc,
    at most k items. Candidate vectors of a mismatched length are dropped.
    """
    q = np.asarray(query, dtype=np.float32)
    dim = q.shape[0]
    texts, vecs = [], []
    for text, vec in candidates:
        if len(vec) == dim:
            texts.append(text)
            vecs.append(vec)
    if not vecs:
        return []
    mat = np.asarray(vecs, dtype=np.float32)
    denom = np.linalg.norm(mat, axis=1) * np.linalg.norm(q)
    denom[denom == 0] = 1e-8
    sims = (mat @ q) / denom
    order = np.argsort(-sims)[:k]
    return [(texts[i], float(sims[i])) for i in order]


def build_retrieval_query(messages: List[Message]) -> str:
    """A short text describing the current interview: opening ask (role/level/
    focus) + the last few turns."""
    users = [m.content for m in messages if m.role == "user"]
    parts: List[str] = []
    if users:
        parts.append(users[0])
    parts.extend(m.content for m in messages[-4:])
    return "\n".join(p for p in parts if p).strip()[:2000]


async def select_related_questions(
    client_id: str, query_text: str, k: int
) -> Optional[List[str]]:
    """Top-k past questions most semantically similar to the current interview.

    Returns None when semantic retrieval isn't usable (embeddings disabled, too
    little embedded history, or the query embedding failed) so the caller can
    fall back to the recency-based list.
    """
    if not EMBEDDINGS_ENABLED:
        return None
    candidates = await fetch_question_embeddings(client_id)
    if len(candidates) < RETRIEVAL_MIN_EMBEDDED:
        return None
    qvec = await embed_text(query_text)
    if qvec is None:
        return None
    return [text for text, _score in cosine_top_k(qvec, candidates, k)]


def build_dedup_block(asked: List[dict], related: Optional[List[str]] = None) -> str:
    """Render the 'do not repeat these' instruction appended to the system prompt.

    `related`: for a large history, the questions semantically closest to the
    current interview (from `select_related_questions`). When given, they replace
    the "most recent N" slice; everything else (category tally, wording) is
    unchanged. When None, the recency slice is used (old behaviour / fallback).
    """
    total = len(asked)
    if total == 0:
        return ""

    header = (
        "\n\n---\n"
        "## ⚠️ 跨场次去重要求（必须遵守）\n"
        "该候选人在**之前的其它面试场次**中已经被问过下面这些题目。"
        "本场面试**禁止**再次提出与它们**重复或高度相似**的题目"
        "（判断标准：同一个考点 / 换个说法的同一道题 / 只改了数字或场景包装）。"
        "请提出**全新的、覆盖不同知识点**的题目。\n\n"
    )

    if total <= DEDUP_FULL_LIMIT:
        body = "\n".join(f"- {a['question_text']}" for a in asked)
        return header + f"已问过的全部 {total} 道题目：\n{body}\n"

    # Large history: category distribution (unchanged) + a bounded list of past
    # questions — semantically-related ones if we have them, else the recent slice.
    cats: dict[str, int] = {}
    for a in asked:
        key = (a["category"] or "未分类").strip() or "未分类"
        cats[key] = cats.get(key, 0) + 1
    tally = "、".join(
        f"{k} {v} 道" for k, v in sorted(cats.items(), key=lambda kv: -kv[1])
    )

    if related:
        listed = related
        intro = (
            f"历史累计已问 {total} 道题，按类别分布：{tally}。\n"
            f"下面是与**本场面试方向语义最相关**的 {len(listed)} 道历史题目"
            f"（务必避开这些及其相似题），并请优先覆盖上面分布中占比更少的类别：\n"
        )
    else:
        listed = [a["question_text"] for a in asked[:DEDUP_RECENT_N]]
        intro = (
            f"历史累计已问 {total} 道题，按类别分布：{tally}。\n"
            f"其中最近 {len(listed)} 道题目如下（务必避开这些及其相似题），"
            f"并请优先覆盖上面分布中占比更少的类别：\n"
        )
    return header + intro + "\n".join(f"- {t}" for t in listed) + "\n"


async def record_asked_questions(
    client: AsyncOpenAI,
    model: str,
    client_id: str,
    session_id: Optional[str],
    assistant_text: str,
) -> None:
    """Extract the interview question(s) from an interviewer turn and persist them.

    Best-effort: any failure (bad JSON, provider without json mode, etc.) is
    swallowed so it can never break the chat response.
    """
    if not client_id or not assistant_text.strip():
        return
    try:
        resp = await client.chat.completions.create(
            model=model,
            max_tokens=800,
            messages=[
                {"role": "system", "content": EXTRACT_PROMPT},
                {"role": "user", "content": assistant_text},
            ],
            response_format={"type": "json_object"},
        )
        data = json.loads(resp.choices[0].message.content)
        questions = data.get("questions") or []
    except Exception:
        return

    clean: List[Tuple[str, Optional[str]]] = []
    for q in questions:
        if not isinstance(q, dict):
            continue
        text = (q.get("text") or "").strip()
        if not text:
            continue
        category = (q.get("category") or "").strip() or None
        clean.append((text, category))

    if not clean:
        return

    async with aiosqlite.connect(DB_PATH) as db:
        new_ids: List[Tuple[int, str]] = []
        for text, category in clean:
            cur = await db.execute(
                "INSERT INTO asked_questions "
                "(client_id, question_text, category, session_id) VALUES (?, ?, ?, ?)",
                (client_id, text, category, session_id),
            )
            new_ids.append((cur.lastrowid, text))
        await db.commit()

        # Best-effort: embed each new question and backfill the vector. Any
        # failure here leaves embedding NULL — the question is already recorded.
        try:
            for qid, text in new_ids:
                vec = await embed_text(text)
                if vec is None:
                    continue
                await db.execute(
                    "UPDATE asked_questions SET embedding = ? WHERE id = ?",
                    (json.dumps(vec), qid),
                )
            await db.commit()
        except Exception:
            pass


# ── Cross-session candidate profile ──────────────
async def fetch_client_profile(client_id: str) -> str:
    """The stored free-text profile for this client ('' if none)."""
    if not client_id:
        return ""
    async with aiosqlite.connect(DB_PATH) as db:
        async with db.execute(
            "SELECT profile FROM client_profile WHERE client_id = ?", (client_id,)
        ) as cur:
            row = await cur.fetchone()
    return (row[0] if row else "") or ""


def build_profile_block(profile: str) -> str:
    """The 'here's who the candidate is' section appended to the system prompt."""
    p = (profile or "").strip()
    if not p:
        return ""
    return (
        "\n\n---\n"
        "## 该候选人的已知信息（跨场次记忆，请直接使用）\n"
        f"{p}\n\n"
        "请在本场面试中自然地运用上述信息（例如用候选人的姓名称呼对方），"
        "不要再重复询问这里已经写明的内容。\n"
        "注意：这些是**历史场次**攒下的记忆，可能已过时。"
        "如果与上面「本场面试的定向出题材料」里的简历 / JD 有冲突"
        "（例如项目经历、技术栈对不上），**一律以本场简历 / JD 为准**，"
        "不要提问这里写了但简历里没有的项目。\n"
    )


# ── Résumé + JD (定向出题) ───────────────────────
def _extract_resume_text(filename: str, raw: bytes) -> str:
    """Pull plain text out of an uploaded résumé (pdf / docx / txt)."""
    ext = Path(filename or "").suffix.lower()
    if ext == ".pdf":
        import pdfplumber  # heavy import — only when a PDF is actually uploaded
        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    elif ext == ".docx":
        doc = Document(io.BytesIO(raw))
        text = "\n".join(p.text for p in doc.paragraphs)
    elif ext in (".txt", ".md"):
        text = raw.decode("utf-8", errors="ignore")
    else:
        raise HTTPException(400, "只支持 PDF、Word（.docx）或 txt 格式的简历")
    return text.strip()


async def fetch_resume(client_id: str) -> dict:
    """This client's stored résumé row ({} if none)."""
    if not client_id:
        return {}
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            "SELECT resume_text, resume_filename, updated_at FROM user_profile "
            "WHERE client_id = ?",
            (client_id,),
        ) as cur:
            row = await cur.fetchone()
    return dict(row) if row else {}


async def fetch_jd(client_id: str, jd_id: Optional[int]) -> Optional[dict]:
    """One JD, scoped to this client (None if it doesn't exist / isn't theirs)."""
    if not client_id or not jd_id:
        return None
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            "SELECT id, title, jd_text FROM job_descriptions "
            "WHERE id = ? AND client_id = ?",
            (jd_id, client_id),
        ) as cur:
            row = await cur.fetchone()
    return dict(row) if row else None


def build_jobfit_block(resume_text: str, jd: Optional[dict]) -> str:
    """The 'tailor the questions to this résumé + JD' section of the prompt."""
    resume_text = (resume_text or "").strip()[:RESUME_MAX_CHARS]
    jd_text = ((jd or {}).get("jd_text") or "").strip()[:JD_MAX_CHARS]
    if not resume_text and not jd_text:
        return ""
    out = [
        "\n\n---\n"
        "## 本场面试的定向出题材料（优先级高于通用出题）\n"
        "请**优先围绕下面的简历项目经历和岗位技术栈出题**："
        "先深挖简历里的真实项目（追问技术选型、遇到的难点、你的具体职责与权衡），"
        "再结合 JD 要求的技术栈补充针对性考题。"
        "简历 / JD 未提及的方向可以少量涉及，但不要明显偏离。\n"
    ]
    if resume_text:
        out.append(f"\n### 候选人简历\n{resume_text}\n")
    if jd_text:
        title = ((jd or {}).get("title") or "").strip()
        suffix = f"（{title}）" if title else ""
        out.append(f"\n### 目标岗位 JD{suffix}\n{jd_text}\n")
    return "".join(out)


def _conversation_excerpt(messages: List[Message], keep: int = 8) -> str:
    tail = messages[-keep:]
    role = {"user": "候选人", "assistant": "面试官"}
    return "\n".join(f"{role.get(m.role, m.role)}：{m.content}" for m in tail)[:4000]


async def update_client_profile(
    client: AsyncOpenAI, model: str, client_id: str, messages: List[Message]
) -> None:
    """Re-derive the candidate profile from (stored profile + latest turns) and
    upsert it. Best-effort — never raises into the chat response."""
    if not client_id or not messages:
        return
    current = await fetch_client_profile(client_id)
    excerpt = _conversation_excerpt(messages)
    if not excerpt.strip():
        return
    try:
        resp = await client.chat.completions.create(
            model=model,
            max_tokens=700,
            messages=[
                {"role": "system", "content": PROFILE_PROMPT},
                {
                    "role": "user",
                    "content": f"【当前档案】\n{current or '(空)'}\n\n【最新对话片段】\n{excerpt}",
                },
            ],
            response_format={"type": "json_object"},
        )
        new_profile = (json.loads(resp.choices[0].message.content).get("profile") or "").strip()
    except Exception:
        return

    new_profile = new_profile[:PROFILE_MAX_CHARS]
    if not new_profile or new_profile == (current or "").strip():
        return
    try:
        async with aiosqlite.connect(DB_PATH) as db:
            await db.execute(
                "INSERT INTO client_profile (client_id, profile) VALUES (?, ?) "
                "ON CONFLICT(client_id) DO UPDATE SET "
                "profile = excluded.profile, updated_at = datetime('now','localtime')",
                (client_id, new_profile),
            )
            await db.commit()
    except Exception:
        pass


def build_docx(report: dict) -> io.BytesIO:
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2.5)
        sec.left_margin = sec.right_margin = Cm(3)

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("面  试  报  告")
    r.font.size = Pt(22)
    r.font.bold = True

    doc.add_paragraph()

    # Meta
    for line in [
        f"生成时间：{datetime.now().strftime('%Y年%m月%d日 %H:%M')}",
        f"面试岗位：{report.get('role','—')}    经验层级：{report.get('level','—')}",
        f"综合评分：{report.get('overall_score','—')}/10    面试结论：{report.get('verdict','—')}",
    ]:
        doc.add_paragraph(line).runs[0].font.size = Pt(11)

    doc.add_paragraph()

    # Questions
    doc.add_heading("详细题目分析", 1)
    for i, q in enumerate(report.get("questions", []), 1):
        doc.add_heading(f"Q{i}.  {q.get('question','')}", 2)

        # Score
        p = doc.add_paragraph()
        r = p.add_run(f"得分：{q.get('score','—')}/10")
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x21, 0x7D, 0xBB)

        # User answer
        p = doc.add_paragraph()
        r = p.add_run("候选人回答：\n")
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x6B, 0x6B, 0x6B)
        p.add_run(q.get("user_answer", "未作答")).font.size = Pt(10.5)

        # Standard answer
        p = doc.add_paragraph()
        r = p.add_run("标准答案：\n")
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x27, 0x9E, 0x6E)
        p.add_run(q.get("standard_answer", "")).font.size = Pt(10.5)

        # Feedback
        if q.get("feedback"):
            p = doc.add_paragraph()
            r = p.add_run("评价：")
            r.font.bold = True
            r.font.color.rgb = RGBColor(0x8B, 0x5C, 0xF6)
            p.add_run(q["feedback"]).font.size = Pt(10.5)

        doc.add_paragraph()

    # Summary
    doc.add_heading("综合评价", 1)
    for label, key in [
        ("主要优势", "strengths"),
        ("改进方向", "improvements"),
        ("建议学习", "recommended_topics"),
    ]:
        items = report.get(key, [])
        if items:
            doc.add_heading(label, 2)
            for item in items:
                doc.add_paragraph(item, style="List Bullet").runs[0].font.size = Pt(10.5)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ── Health ───────────────────────────────────────
@app.get("/api/health")
async def health():
    return {"status": "ok"}


# ── Sessions ─────────────────────────────────────
@app.post("/api/sessions")
async def create_session():
    sid = str(uuid.uuid4())
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute("INSERT INTO sessions (id) VALUES (?)", (sid,))
        await db.commit()
    return {"id": sid}


@app.get("/api/sessions")
async def list_sessions():
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            "SELECT id, title, turn_count, duration_s, ended, created_at FROM sessions "
            "ORDER BY updated_at DESC LIMIT 60"
        ) as cur:
            rows = await cur.fetchall()
    return [dict(r) for r in rows]


@app.get("/api/sessions/{sid}/messages")
async def get_session_messages(sid: str):
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            "SELECT turn_count, duration_s, ended FROM sessions WHERE id=?", (sid,)
        ) as cur:
            meta = await cur.fetchone()
            if not meta:
                raise HTTPException(404, "Session not found")
        async with db.execute(
            "SELECT role, content FROM messages WHERE session_id=? ORDER BY id", (sid,)
        ) as cur:
            rows = await cur.fetchall()
    return {
        "messages": [dict(r) for r in rows],
        "turn_count": meta["turn_count"],
        "duration_s": meta["duration_s"],
        "ended": meta["ended"],
    }


@app.post("/api/sessions/save-turn")
async def save_turn(data: TurnSave):
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute(
            "INSERT INTO messages (session_id, role, content) VALUES (?,?,?)",
            (data.session_id, "user", data.user_content),
        )
        await db.execute(
            "INSERT INTO messages (session_id, role, content) VALUES (?,?,?)",
            (data.session_id, "assistant", data.assistant_content),
        )
        if data.title:
            await db.execute(
                "UPDATE sessions SET title=?,turn_count=?,duration_s=?,"
                "updated_at=datetime('now','localtime') WHERE id=?",
                (data.title, data.turn_count, data.duration_s, data.session_id),
            )
        else:
            await db.execute(
                "UPDATE sessions SET turn_count=?,duration_s=?,"
                "updated_at=datetime('now','localtime') WHERE id=?",
                (data.turn_count, data.duration_s, data.session_id),
            )
        await db.commit()
    return {"ok": True}


@app.post("/api/sessions/{sid}/end")
async def end_session(sid: str):
    """Mark an interview as finished (candidate sent `end` / clicked 结束面试).
    The frontend calls this the moment the interview ends, so the sidebar can
    show 已通关 and a reload / session-switch restores the ended state."""
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute(
            "UPDATE sessions SET ended=1, updated_at=datetime('now','localtime') WHERE id=?",
            (sid,),
        )
        await db.commit()
    return {"ok": True}


@app.delete("/api/sessions/{sid}")
async def delete_session(sid: str):
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute("DELETE FROM messages WHERE session_id=?", (sid,))
        await db.execute("DELETE FROM sessions WHERE id=?", (sid,))
        await db.commit()
    return {"ok": True}


# ── Report ───────────────────────────────────────
@app.post("/api/sessions/{sid}/report")
async def generate_report(sid: str, body: ReportRequest = None):
    # Load conversation
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            "SELECT role, content FROM messages WHERE session_id=? ORDER BY id", (sid,)
        ) as cur:
            rows = await cur.fetchall()

    msgs = [dict(r) for r in rows]
    if len(msgs) < 4:
        raise HTTPException(400, "对话内容不足，无法生成报告")

    conv = "\n\n".join(
        f"{'面试官' if m['role']=='assistant' else '候选人'}：{m['content']}"
        for m in msgs
    )

    client, model = make_client(body.llm if body else None)

    resp = await client.chat.completions.create(
        model=model,
        max_tokens=4096,
        messages=[
            {"role": "system", "content": REPORT_PROMPT},
            {"role": "user", "content": f"面试对话如下：\n\n{conv}"},
        ],
        response_format={"type": "json_object"},
    )

    report = json.loads(resp.choices[0].message.content)

    # Persist so download doesn't need to regenerate. Generating a report also
    # means the interview is over — mark it ended in case `end` was never sent.
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute(
            "UPDATE sessions SET report_json=?, ended=1 WHERE id=?",
            (json.dumps(report, ensure_ascii=False), sid),
        )
        await db.commit()

    return report


@app.get("/api/sessions/{sid}/report/download")
async def download_report(sid: str):
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            "SELECT report_json FROM sessions WHERE id=?", (sid,)
        ) as cur:
            row = await cur.fetchone()

    if not row or not row["report_json"]:
        raise HTTPException(404, "请先生成报告")

    report = json.loads(row["report_json"])
    buf = build_docx(report)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''interview_report.docx"},
    )


# ── Chat ─────────────────────────────────────────
def _last_user_text(messages: List[Message]) -> str:
    for m in reversed(messages):
        if m.role == "user":
            return m.content.strip()
    return ""


def _is_command_turn(text: str) -> bool:
    if not text:
        return False
    first = text.lower().split()[0] if text.split() else ""
    return first in INTERVIEW_COMMANDS


# Detached background jobs (post-turn question extraction / profile refresh). We
# keep hard references so the event loop doesn't GC a running task.
_bg_tasks: set = set()


def _run_detached(coro) -> None:
    task = asyncio.create_task(coro)
    _bg_tasks.add(task)
    task.add_done_callback(_bg_tasks.discard)


async def _post_turn(
    client, model, client_id, session_id, assistant_text, messages, skip_extract=False
) -> None:
    """Fire-and-forget after a turn is streamed: record the interviewer's new
    questions and refresh the candidate profile. Runs *after* the SSE stream has
    already closed, so nothing here can delay or lose the client's turn save.

    skip_extract: set on 'replay' turns — the asked question is already known and
    was recorded synchronously, so there's nothing to extract."""
    if not skip_extract:
        try:
            await record_asked_questions(client, model, client_id, session_id, assistant_text)
        except Exception:
            pass
    try:
        await update_client_profile(client, model, client_id, messages)
    except Exception:
        pass


# ── 面试知识库 ───────────────────────────────────
def _kb_extract_text(filename: str, raw: bytes) -> str:
    """Plain text out of an uploaded knowledge-base file (.md / .docx / .txt)."""
    ext = Path(filename or "").suffix.lower()
    if ext in (".md", ".markdown", ".txt"):
        return raw.decode("utf-8", errors="ignore").strip()
    if ext == ".docx":
        doc = Document(io.BytesIO(raw))
        return "\n".join(p.text for p in doc.paragraphs).strip()
    raise HTTPException(400, "只支持 .md 或 .docx 文件")


async def _kb_count() -> int:
    async with aiosqlite.connect(DB_PATH) as db:
        async with db.execute("SELECT COUNT(*) FROM question_bank") as cur:
            return (await cur.fetchone())[0]


async def _kb_pick_unasked(client_id: str) -> Optional[dict]:
    """A random bank question this client has never been asked (by exact text)."""
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            "SELECT id, question_text, category FROM question_bank "
            "WHERE question_text NOT IN "
            "  (SELECT question_text FROM asked_questions WHERE client_id = ?) "
            "ORDER BY RANDOM() LIMIT 1",
            (client_id,),
        ) as cur:
            row = await cur.fetchone()
    return dict(row) if row else None


async def _kb_style_examples(n: int = KB_STYLE_EXAMPLE_N) -> List[dict]:
    """A handful of bank questions to few-shot the style branch. Prefers a single
    category (picked at random from one that has enough rows) so the examples
    hang together; falls back to a plain random sample."""
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            "SELECT category FROM question_bank WHERE category IS NOT NULL "
            "GROUP BY category HAVING COUNT(*) >= 3 ORDER BY RANDOM() LIMIT 1"
        ) as cur:
            cat_row = await cur.fetchone()
        if cat_row:
            async with db.execute(
                "SELECT question_text, category FROM question_bank WHERE category = ? "
                "ORDER BY RANDOM() LIMIT ?",
                (cat_row["category"], n),
            ) as cur:
                rows = await cur.fetchall()
        else:
            async with db.execute(
                "SELECT question_text, category FROM question_bank ORDER BY RANDOM() LIMIT ?",
                (n,),
            ) as cur:
                rows = await cur.fetchall()
    return [dict(r) for r in rows]


async def build_question_bank_block(client_id: str, messages: List[Message]):
    """Decide how the next interview question is produced and return
    (prompt_block, replay_question_or_None).

    Only kicks in once the interview proper has started (there's at least one
    interviewer turn already) and only when the bank is non-empty — otherwise
    returns ('', None) and the existing free-generation logic is untouched."""
    if not any(m.role == "assistant" for m in messages):
        return "", None
    if await _kb_count() == 0:
        return "", None

    roll = random.random()

    # a) 原题重放
    if roll < KB_REPLAY_PROB:
        q = await _kb_pick_unasked(client_id)
        if q:
            block = (
                "\n\n---\n"
                "## 🎯 本轮出题指令（最高优先级，必须严格执行）\n"
                "先简短点评候选人上一题的回答，然后**一字不差地**向候选人提出下面这道题，"
                "不要改写、不要合并、不要加编号前缀，把它作为本轮的正式面试题：\n\n"
                f"「{q['question_text']}」\n"
            )
            return block, q["question_text"]
        roll = KB_REPLAY_PROB  # nothing left to replay → fall into the style range

    # b) 风格生成
    if roll < KB_REPLAY_PROB + KB_STYLE_PROB:
        examples = await _kb_style_examples()
        if examples:
            listed = "\n".join(
                f"- 【{e['category'] or '未分类'}】{e['question_text']}" for e in examples
            )
            block = (
                "\n\n---\n"
                "## 本轮出题参考（候选人整理的真实面试题）\n"
                "下面是候选人自己遇到过的真实面试题。请**模仿它们的提问风格、难度、"
                "切入角度和考察点**，出**一道全新的**同类型问题——不要照抄、不要只换"
                "数字或措辞，要覆盖示例里没问到的知识点：\n\n"
                f"{listed}\n"
            )
            return block, None

    # c) 自由生成 —— 不加任何 block
    return "", None


async def _embed_backfill(qid: int, text: str) -> None:
    try:
        vec = await embed_text(text)
        if vec is None:
            return
        async with aiosqlite.connect(DB_PATH) as db:
            await db.execute(
                "UPDATE asked_questions SET embedding = ? WHERE id = ?", (json.dumps(vec), qid)
            )
            await db.commit()
    except Exception:
        pass


async def record_replayed_question(client_id: str, session_id: Optional[str], question_text: str) -> None:
    """Record a verbatim-replayed bank question into asked_questions, tagged so
    it's distinguishable from LLM-generated ones. The INSERT is synchronous (the
    turn must not repeat this question later); the embedding is backfilled off
    the hot path."""
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            "SELECT category FROM question_bank WHERE question_text = ? LIMIT 1",
            (question_text,),
        ) as cur:
            row = await cur.fetchone()
        category = row["category"] if row else None
        cur = await db.execute(
            "INSERT INTO asked_questions "
            "(client_id, question_text, category, session_id, source) VALUES (?, ?, ?, ?, ?)",
            (client_id, question_text, category, session_id, "knowledge_base"),
        )
        await db.commit()
        qid = cur.lastrowid
    _run_detached(_embed_backfill(qid, question_text))


@app.post("/api/chat")
async def chat(request: ChatRequest, x_client_id: Optional[str] = Header(default=None)):
    client, model = make_client(request.llm)

    # 1) Pull this client's full question history and build the avoidance block.
    #    For a large history, pick the semantically-related past questions
    #    (embeddings + in-memory cosine) instead of just the most recent ones.
    asked = await fetch_asked_questions(x_client_id) if x_client_id else []
    related = None
    if x_client_id and len(asked) > DEDUP_FULL_LIMIT:
        related = await select_related_questions(
            x_client_id, build_retrieval_query(request.messages), RETRIEVAL_TOP_K
        )

    # 2) Load the cross-session candidate profile (name, target role, ...).
    profile = await fetch_client_profile(x_client_id) if x_client_id else ""

    # 3) Optional job-fit material: the candidate's résumé + a chosen JD. When
    #    neither is selected this contributes nothing and the interview runs in
    #    the existing generic mode.
    resume = await fetch_resume(x_client_id) if (x_client_id and request.use_resume) else {}
    jd = await fetch_jd(x_client_id, request.jd_id) if x_client_id else None

    # Command turns (hint/skip/explain/...) don't introduce new questions.
    record_new = bool(x_client_id) and not _is_command_turn(_last_user_text(request.messages))

    # 4) Interview knowledge base: maybe replay a real question verbatim, or
    #    few-shot the model with real questions to shape a new one. No-op when
    #    the bank is empty or we roll "free generate".
    kb_block, replay_q = ("", None)
    if record_new:
        kb_block, replay_q = await build_question_bank_block(x_client_id, request.messages)

    # Order matters: the résumé/JD block goes *before* the cross-session profile
    # so a freshly uploaded résumé takes precedence over stale remembered facts
    # (see build_profile_block's "以本场简历 / JD 为准" note).
    system_prompt = (
        SYSTEM_PROMPT
        + build_jobfit_block(resume.get("resume_text", ""), jd)
        + build_profile_block(profile)
        + build_dedup_block(asked, related)
        + kb_block
    )

    if replay_q:
        await record_replayed_question(x_client_id, request.session_id, replay_q)

    async def generate():
        full = ""
        try:
            stream = await client.chat.completions.create(
                model=model,
                max_tokens=2048,
                messages=[
                    {"role": "system", "content": system_prompt},
                    *[{"role": m.role, "content": m.content} for m in request.messages],
                ],
                stream=True,
            )
            async for chunk in stream:
                text = chunk.choices[0].delta.content
                if text:
                    full += text
                    yield f"data: {json.dumps({'type':'text','content':text}, ensure_ascii=False)}\n\n"

            # Tell the client the answer is complete *first* — it saves the turn
            # on this event. The cross-session bookkeeping below makes its own
            # extra LLM calls (slow), so it must NOT sit between the last token
            # and 'done', or a refresh in that window loses the whole turn.
            yield f"data: {json.dumps({'type':'done'})}\n\n"

            # 4) Persist what the interviewer just asked (so future sessions
            #    don't repeat it) and refresh the candidate profile — detached,
            #    after the stream has closed.
            if record_new and full.strip():
                _run_detached(_post_turn(
                    client, model, x_client_id, request.session_id, full, request.messages,
                    skip_extract=bool(replay_q),
                ))
        except Exception as e:
            yield f"data: {json.dumps({'type':'error','content':str(e)}, ensure_ascii=False)}\n\n"

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no", "Connection": "keep-alive"},
    )


# ── Asked-questions inspection / reset (handy for testing) ──
@app.get("/api/asked-questions")
async def list_asked_questions(x_client_id: Optional[str] = Header(default=None)):
    if not x_client_id:
        return {"client_id": None, "count": 0, "embedded": 0, "questions": []}
    rows = await fetch_asked_questions(x_client_id)
    embedded = len(await fetch_question_embeddings(x_client_id))
    return {
        "client_id": x_client_id,
        "count": len(rows),
        "embedded": embedded,
        "embeddings_enabled": EMBEDDINGS_ENABLED,
        "questions": rows,
    }


@app.get("/api/asked-questions/search")
async def search_asked_questions(
    q: str, k: int = RETRIEVAL_TOP_K, x_client_id: Optional[str] = Header(default=None)
):
    """Debug: run the same semantic retrieval the chat endpoint uses and return
    the ranked past questions with scores, so you can eyeball relevance."""
    if not x_client_id:
        raise HTTPException(400, "缺少 X-Client-Id 请求头")
    if not EMBEDDINGS_ENABLED:
        return {"enabled": False, "results": [], "detail": "EMBEDDING_API_KEY 未配置"}
    candidates = await fetch_question_embeddings(x_client_id)
    qvec = await embed_text(q)
    if qvec is None:
        return {"enabled": True, "results": [], "detail": "query embedding 失败"}
    ranked = cosine_top_k(qvec, candidates, k)
    return {
        "enabled": True,
        "query": q,
        "pool": len(candidates),
        "results": [{"score": round(s, 4), "question": t} for t, s in ranked],
    }


@app.delete("/api/asked-questions")
async def clear_asked_questions(x_client_id: Optional[str] = Header(default=None)):
    if not x_client_id:
        raise HTTPException(400, "缺少 X-Client-Id 请求头")
    async with aiosqlite.connect(DB_PATH) as db:
        cur = await db.execute(
            "DELETE FROM asked_questions WHERE client_id = ?", (x_client_id,)
        )
        await db.commit()
        deleted = cur.rowcount
    return {"ok": True, "deleted": deleted}


# ── Candidate profile inspection / reset (handy for testing) ──
@app.get("/api/profile")
async def get_profile(x_client_id: Optional[str] = Header(default=None)):
    if not x_client_id:
        return {"client_id": None, "profile": ""}
    return {"client_id": x_client_id, "profile": await fetch_client_profile(x_client_id)}


@app.delete("/api/profile")
async def clear_profile(x_client_id: Optional[str] = Header(default=None)):
    if not x_client_id:
        raise HTTPException(400, "缺少 X-Client-Id 请求头")
    async with aiosqlite.connect(DB_PATH) as db:
        cur = await db.execute(
            "DELETE FROM client_profile WHERE client_id = ?", (x_client_id,)
        )
        await db.commit()
        deleted = cur.rowcount
    return {"ok": True, "deleted": deleted}


# ── Résumé + JD (定向出题) ───────────────────────
@app.get("/api/resume")
async def get_resume(x_client_id: Optional[str] = Header(default=None)):
    row = await fetch_resume(x_client_id) if x_client_id else {}
    text = row.get("resume_text", "")
    return {
        "filename": row.get("resume_filename"),
        "resume_text": text,
        "chars": len(text),
        "updated_at": row.get("updated_at"),
    }


@app.post("/api/resume")
async def upload_resume(
    file: UploadFile = File(...),
    x_client_id: Optional[str] = Header(default=None),
):
    if not x_client_id:
        raise HTTPException(400, "缺少 X-Client-Id 请求头")
    raw = await file.read()
    if not raw:
        raise HTTPException(400, "这个文件是空的，换一份试试？")
    text = _extract_resume_text(file.filename, raw)[:RESUME_MAX_CHARS]
    if not text:
        raise HTTPException(
            400, "没能从文件里读到文字，可能是扫描版 PDF —— 导出成文本或 Word 再试试"
        )
    async with aiosqlite.connect(DB_PATH) as db:
        await db.execute(
            "INSERT INTO user_profile (client_id, resume_text, resume_filename) "
            "VALUES (?, ?, ?) "
            "ON CONFLICT(client_id) DO UPDATE SET "
            "resume_text = excluded.resume_text, "
            "resume_filename = excluded.resume_filename, "
            "updated_at = datetime('now','localtime')",
            (x_client_id, text, file.filename),
        )
        # A new résumé means "get to know this candidate again": the LLM-maintained
        # cross-session profile is merge-only and never drops a stale project, so
        # wipe it here. It rebuilds from the next interview's turns.
        await db.execute(
            "DELETE FROM client_profile WHERE client_id = ?", (x_client_id,)
        )
        await db.commit()
    return {"ok": True, "filename": file.filename, "chars": len(text)}


@app.delete("/api/resume")
async def delete_resume(x_client_id: Optional[str] = Header(default=None)):
    if not x_client_id:
        raise HTTPException(400, "缺少 X-Client-Id 请求头")
    async with aiosqlite.connect(DB_PATH) as db:
        cur = await db.execute(
            "DELETE FROM user_profile WHERE client_id = ?", (x_client_id,)
        )
        await db.commit()
    return {"ok": True, "deleted": cur.rowcount}


@app.get("/api/jds")
async def list_jds(x_client_id: Optional[str] = Header(default=None)):
    if not x_client_id:
        return []
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            "SELECT id, title, jd_text, created_at FROM job_descriptions "
            "WHERE client_id = ? ORDER BY id DESC",
            (x_client_id,),
        ) as cur:
            rows = await cur.fetchall()
    return [dict(r) for r in rows]


@app.post("/api/jds")
async def create_jd(data: JDCreate, x_client_id: Optional[str] = Header(default=None)):
    if not x_client_id:
        raise HTTPException(400, "缺少 X-Client-Id 请求头")
    jd_text = data.jd_text.strip()[:JD_MAX_CHARS]
    if not jd_text:
        raise HTTPException(400, "JD 内容不能为空")
    title = (data.title or "").strip()[:120] or "未命名岗位"
    async with aiosqlite.connect(DB_PATH) as db:
        cur = await db.execute(
            "INSERT INTO job_descriptions (client_id, title, jd_text) VALUES (?, ?, ?)",
            (x_client_id, title, jd_text),
        )
        await db.commit()
        new_id = cur.lastrowid
    return {"id": new_id, "title": title}


@app.delete("/api/jds/{jd_id}")
async def delete_jd(jd_id: int, x_client_id: Optional[str] = Header(default=None)):
    if not x_client_id:
        raise HTTPException(400, "缺少 X-Client-Id 请求头")
    async with aiosqlite.connect(DB_PATH) as db:
        cur = await db.execute(
            "DELETE FROM job_descriptions WHERE id = ? AND client_id = ?",
            (jd_id, x_client_id),
        )
        await db.commit()
    return {"ok": True, "deleted": cur.rowcount}


# ── 面试知识库 ───────────────────────────────────
@app.post("/api/knowledge-base/import")
async def import_knowledge_base(
    file: UploadFile = File(...),
    api_key: Optional[str] = Form(None),
    base_url: Optional[str] = Form(None),
    model: Optional[str] = Form(None),
):
    raw = await file.read()
    if not raw:
        raise HTTPException(400, "这个文件是空的")
    text = _kb_extract_text(file.filename, raw)
    if not text:
        raise HTTPException(400, "没能从文件里读到文字")
    text = text[:KB_MAX_IMPORT_CHARS]

    cfg = LLMConfig(api_key=api_key, base_url=base_url, model=model) if (api_key and base_url and model) else None
    client, model_name = make_client(cfg)

    items: List[dict] = []
    try:
        resp = await client.chat.completions.create(
            model=model_name,
            max_tokens=4000,
            messages=[
                {"role": "system", "content": KB_EXTRACT_PROMPT},
                {"role": "user", "content": text},
            ],
            response_format={"type": "json_object"},
        )
        data = json.loads(resp.choices[0].message.content)
        items = data.get("questions") or []
    except Exception:
        items = []

    rows: List[Tuple[str, Optional[str], Optional[str]]] = []
    for it in items:
        if not isinstance(it, dict):
            continue
        q = (it.get("question") or it.get("text") or "").strip()
        if not q:
            continue
        cat = (it.get("category") or "").strip() or None
        cor = (it.get("company_or_role") or "").strip() or None
        rows.append((q, cat, cor))

    fallback = not rows
    async with aiosqlite.connect(DB_PATH) as db:
        if rows:
            for q, cat, cor in rows:
                await db.execute(
                    "INSERT INTO question_bank (question_text, category, source_file, company_or_role) "
                    "VALUES (?, ?, ?, ?)",
                    (q, cat, file.filename, cor),
                )
        else:
            # Extraction failed — keep the whole file as one bucket row so the
            # content isn't lost; it can still be inspected / deleted.
            await db.execute(
                "INSERT INTO question_bank (question_text, category, source_file, company_or_role) "
                "VALUES (?, NULL, ?, NULL)",
                (text, file.filename),
            )
        await db.commit()

    return {"ok": True, "imported": len(rows) if rows else 1, "fallback": fallback, "source_file": file.filename}


@app.get("/api/knowledge-base")
async def list_knowledge_base(
    category: Optional[str] = None, offset: int = 0, limit: int = 50
):
    limit = max(1, min(limit, 200))
    offset = max(0, offset)
    where, params = "", []
    if category:
        where = "WHERE category = ?"
        params.append(category)
    async with aiosqlite.connect(DB_PATH) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            f"SELECT COUNT(*) FROM question_bank {where}", params
        ) as cur:
            total = (await cur.fetchone())[0]
        async with db.execute(
            f"SELECT id, question_text, category, source_file, company_or_role, created_at "
            f"FROM question_bank {where} ORDER BY id DESC LIMIT ? OFFSET ?",
            [*params, limit, offset],
        ) as cur:
            rows = await cur.fetchall()
        async with db.execute(
            "SELECT category, COUNT(*) n FROM question_bank "
            "WHERE category IS NOT NULL GROUP BY category ORDER BY category"
        ) as cur:
            cats = [{"category": r[0], "count": r[1]} for r in await cur.fetchall()]
    return {
        "total": total,
        "offset": offset,
        "limit": limit,
        "items": [dict(r) for r in rows],
        "categories": cats,
    }


@app.delete("/api/knowledge-base/{item_id}")
async def delete_knowledge_base(item_id: int):
    async with aiosqlite.connect(DB_PATH) as db:
        cur = await db.execute("DELETE FROM question_bank WHERE id = ?", (item_id,))
        await db.commit()
    return {"ok": True, "deleted": cur.rowcount}


# Must be last
frontend_dir = Path(__file__).parent.parent / "frontend"
app.mount("/", StaticFiles(directory=str(frontend_dir), html=True), name="static")
